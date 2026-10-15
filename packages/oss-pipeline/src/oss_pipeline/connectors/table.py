"""Table connector: each row of a CSV, TSV or XLSX file becomes a document.

Settings: `path` (file or folder), `title_column`, `text_columns` (list; all
columns when empty), `id_column` (row number when empty), `field_columns`
(mapping of column -> suffixed field name such as `owner_s`), `sheet`
(XLSX), `encoding` (utf-8). The row is rendered as "Column: value" lines so
Tika and the patterns stage see readable text; mapped columns become facet
fields through `x_<field>` as the index stage expects.
"""

from __future__ import annotations

import csv
import hashlib
import tempfile
from collections.abc import Iterable, Iterator
from pathlib import Path
from typing import Any

from oss_core.contracts import Cursor, FetchedItem, SourceConfig, SourceItem


def _rows(path: Path, settings: dict[str, Any]) -> Iterator[dict[str, str]]:
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            return
        wb = load_workbook(str(path), read_only=True, data_only=True)
        ws = wb[settings["sheet"]] if settings.get("sheet") in wb.sheetnames else wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        header = [str(h or f"col{i}") for i, h in enumerate(next(rows, []) or [])]
        for values in rows:
            yield {
                header[i]: ("" if v is None else str(v))
                for i, v in enumerate(values)
                if i < len(header)
            }
        return
    delimiter = "\t" if suffix == ".tsv" else ","
    with path.open(encoding=str(settings.get("encoding") or "utf-8"), newline="") as fh:
        reader = csv.DictReader(fh, delimiter=delimiter)
        for row in reader:
            yield {str(k): ("" if v is None else str(v)) for k, v in row.items() if k is not None}


class TableConnector:
    name = "table"

    def discover(self, source: SourceConfig, since: Cursor | None) -> Iterable[SourceItem]:
        root = Path(str(source.settings["path"])).expanduser()
        files = (
            [root]
            if root.is_file()
            else sorted(p for p in root.rglob("*") if p.suffix.lower() in (".csv", ".tsv", ".xlsx"))
        )
        settings = source.settings
        text_columns = [str(c) for c in settings.get("text_columns") or []]
        title_column = str(settings.get("title_column") or "")
        id_column = str(settings.get("id_column") or "")
        field_columns = {str(k): str(v) for k, v in (settings.get("field_columns") or {}).items()}
        tmp = Path(tempfile.gettempdir()) / "semascope-table"
        tmp.mkdir(parents=True, exist_ok=True)
        for file in files:
            for n, row in enumerate(_rows(file, settings), start=1):
                columns = text_columns or list(row)
                lines = [
                    f"{col}: {row.get(col, '')}" for col in columns if str(row.get(col, "")).strip()
                ]
                if not lines:
                    continue
                title = row.get(title_column) if title_column else None
                text = (f"{title}\n\n" if title else "") + "\n".join(lines) + "\n"
                content = text.encode("utf-8")
                digest = hashlib.sha256(content).hexdigest()
                row_id = row.get(id_column) if id_column else str(n)
                out = tmp / f"{digest}.txt"
                out.write_bytes(content)
                fields = {
                    f"x_{field}": row.get(col, "")
                    for col, field in field_columns.items()
                    if row.get(col, "")
                }
                yield SourceItem(
                    uri=f"{file.as_uri()}#row={row_id}",
                    content_hash=digest,
                    size=len(content),
                    metadata={
                        "path": str(out),
                        "media_type": "text/plain",
                        "title": title or f"{file.stem} row {row_id}",
                        "fields": fields,
                    },
                )

    def fetch(self, item: SourceItem) -> FetchedItem:
        return FetchedItem(item=item, path=str(item.metadata["path"]), media_type="text/plain")

    def deleted(self, source: SourceConfig, since: Cursor) -> Iterable[str]:
        return []
